# -*- coding: utf-8 -*-
"""Gera dois arquivos .xlsx (um por fonte) com versao legivel/esquematizada
das tabelas refs_agostini.csv e refs_gonzalez.csv: colunas separadas e
formatacao (cabecalho, largura, congelamento, filtro, quebra de linha).
Nao altera os CSVs originais usados pelo pipeline (consolidar.py etc)."""
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
THIN = Side(style="thin", color="B7C6D6")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(vertical="top", wrap_text=True)
TOP = Alignment(vertical="top")

def formatar_planilha(ws, df, wrap_cols=(), widths=None, freeze="A2"):
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    for col_idx, col_name in enumerate(df.columns, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
        cell.border = BORDER
        letter = get_column_letter(col_idx)
        width = (widths or {}).get(col_name, max(12, min(40, len(str(col_name)) + 4)))
        ws.column_dimensions[letter].width = width
    for row_idx in range(2, ws.max_row + 1):
        for col_idx, col_name in enumerate(df.columns, start=1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = BORDER
            cell.alignment = WRAP if col_name in wrap_cols else TOP
        if row_idx % 2 == 0:
            for col_idx in range(1, len(df.columns) + 1):
                ws.cell(row=row_idx, column=col_idx).fill = PatternFill("solid", fgColor="F2F6FB")
    ws.row_dimensions[1].height = 30

def montar_agostini():
    df = pd.read_csv("dados/refs_agostini.csv", dtype=str).fillna("")
    df = df.rename(columns={
        "ref_id": "ID", "fonte": "Fonte", "autor": "Autor(es)", "ano": "Ano",
        "arxiv": "arXiv", "doi": "DOI", "titulo": "Título",
        "raw": "Citação Completa",
    })
    return df[["ID", "Fonte", "Autor(es)", "Ano", "Título", "arXiv", "DOI", "Citação Completa"]]

def montar_gonzalez():
    df = pd.read_csv("dados/refs_gonzalez.csv", dtype=str).fillna("")
    partes = df["journal"].str.split("|", n=2, expand=True)
    partes = partes.reindex(columns=[0, 1, 2]).fillna("")
    df["Periódico"], df["Volume"], df["Página"] = partes[0], partes[1], partes[2]
    df = df.rename(columns={
        "ref_id": "ID", "fonte": "Fonte", "autor": "Autor(es)", "ano": "Ano",
        "arxiv": "arXiv", "doi": "DOI", "titulo": "Título",
        "raw": "Citação Completa",
    })
    return df[["ID", "Fonte", "Autor(es)", "Ano", "Título", "Periódico", "Volume", "Página",
               "arXiv", "DOI", "Citação Completa"]]

def salvar_arquivo(out_path, sheet_name, df, wrap_cols, widths):
    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name=sheet_name, index=False)
        formatar_planilha(writer.sheets[sheet_name], df, wrap_cols=wrap_cols, widths=widths)
    print(f"OK -> {out_path} ({len(df)} linhas)")

def main():
    df_a = montar_agostini()
    df_g = montar_gonzalez()

    salvar_arquivo(
        "dados/refs_agostini_organizado.xlsx", "Agostini", df_a,
        wrap_cols={"Título", "Citação Completa"},
        widths={"ID": 10, "Fonte": 12, "Autor(es)": 22, "Ano": 8,
                "Título": 55, "arXiv": 16, "DOI": 22, "Citação Completa": 70},
    )
    salvar_arquivo(
        "dados/refs_gonzalez_organizado.xlsx", "Gonzalez-Garcia", df_g,
        wrap_cols={"Título", "Citação Completa"},
        widths={"ID": 10, "Fonte": 18, "Autor(es)": 24, "Ano": 8,
                "Título": 40, "Periódico": 20, "Volume": 10, "Página": 10,
                "arXiv": 16, "DOI": 22, "Citação Completa": 70},
    )

if __name__ == "__main__":
    main()
